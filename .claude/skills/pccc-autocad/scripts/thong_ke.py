"""Thống kê khối lượng PCCC từ một hoặc nhiều bản vẽ, xuất ra Excel.

Đếm block theo tên + layer, cộng chiều dài đường (ống, cáp) theo layer, cộng diện tích
polyline kín theo layer.

Ví dụ:
  python thong_ke.py ban_ve/ --ra khoi_luong.xlsx            # mọi .dxf trong thư mục (ezdxf)
  python thong_ke.py a.dwg b.dwg --ra khoi_luong.xlsx        # .dwg: mở bằng AutoCAD (COM), chỉ đọc
  python thong_ke.py --com-dang-mo --ra khoi_luong.xlsx      # bản vẽ đang mở trong AutoCAD
  python thong_ke.py ban_ve/ --loc-layer PCCC- --ra kl.xlsx  # chỉ lấy layer bắt đầu bằng PCCC-

Giả thiết đơn vị bản vẽ là mm (đổi bằng --don-vi m). Chỉ đọc Model space.
"""
import argparse
import math
from collections import defaultdict
from pathlib import Path


class TongHop:
    def __init__(self):
        self.block = defaultdict(int)      # (file, layer, tên block) -> số lượng
        self.dai = defaultdict(float)      # (file, layer) -> tổng chiều dài (đơn vị bản vẽ)
        self.dien_tich = defaultdict(float)  # (file, layer) -> tổng diện tích
        self.so_vung = defaultdict(int)    # (file, layer) -> số polyline kín


def doc_dxf(duong_dan, th, loc):
    import ezdxf

    doc = ezdxf.readfile(duong_dan)
    ten = Path(duong_dan).name
    for e in doc.modelspace():
        layer = e.dxf.layer
        if loc and not layer.upper().startswith(loc.upper()):
            continue
        loai = e.dxftype()
        if loai == "INSERT":
            th.block[(ten, layer, e.dxf.name)] += 1
        elif loai == "LINE":
            th.dai[(ten, layer)] += math.dist(e.dxf.start, e.dxf.end)
        elif loai == "ARC":
            goc = (e.dxf.end_angle - e.dxf.start_angle) % 360
            th.dai[(ten, layer)] += math.radians(goc) * e.dxf.radius
        elif loai in ("LWPOLYLINE", "POLYLINE"):
            pts = [tuple(p)[:2] for p in (e.get_points("xy") if loai == "LWPOLYLINE" else e.points())]
            if e.is_closed and len(pts) > 2:
                pts_kin = pts + [pts[0]]
                th.dien_tich[(ten, layer)] += abs(sum(
                    pts[i][0] * pts[(i + 1) % len(pts)][1] - pts[(i + 1) % len(pts)][0] * pts[i][1]
                    for i in range(len(pts)))) / 2
                th.so_vung[(ten, layer)] += 1
            else:
                pts_kin = pts
            th.dai[(ten, layer)] += sum(math.dist(pts_kin[i], pts_kin[i + 1]) for i in range(len(pts_kin) - 1))


def doc_com(doc, th, loc):
    import acad_com as ac

    ten = ac.goi_lai(lambda: doc.Name)
    ms = ac.goi_lai(lambda: doc.ModelSpace)
    for i in range(ac.goi_lai(lambda: ms.Count)):
        e = ac.goi_lai(lambda: ms.Item(i))
        layer = ac.goi_lai(lambda: e.Layer)
        if loc and not layer.upper().startswith(loc.upper()):
            continue
        loai = ac.goi_lai(lambda: e.ObjectName)
        if loai == "AcDbBlockReference":
            th.block[(ten, layer, ac.goi_lai(lambda: e.EffectiveName))] += 1
        elif loai in ("AcDbLine", "AcDbPolyline", "AcDb2dPolyline"):
            th.dai[(ten, layer)] += ac.goi_lai(lambda: e.Length)
            if loai != "AcDbLine" and ac.goi_lai(lambda: e.Closed):
                th.dien_tich[(ten, layer)] += ac.goi_lai(lambda: e.Area)
                th.so_vung[(ten, layer)] += 1
        elif loai == "AcDbArc":
            th.dai[(ten, layer)] += ac.goi_lai(lambda: e.ArcLength)


def xuat_excel(th, duong_dan, he_so):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Thiet bi (block)"
    ws.append(["Bản vẽ", "Layer", "Tên block", "Số lượng"])
    for (f, l, b), n in sorted(th.block.items()):
        ws.append([f, l, b, n])

    ws2 = wb.create_sheet("Chieu dai")
    ws2.append(["Bản vẽ", "Layer", "Tổng chiều dài (m)"])
    for (f, l), d in sorted(th.dai.items()):
        ws2.append([f, l, round(d / he_so, 2)])

    ws3 = wb.create_sheet("Dien tich")
    ws3.append(["Bản vẽ", "Layer", "Số vùng kín", "Tổng diện tích (m²)"])
    for (f, l), a in sorted(th.dien_tich.items()):
        ws3.append([f, l, th.so_vung[(f, l)], round(a / he_so ** 2, 2)])

    # Bảng gộp mọi bản vẽ theo tên block, tiện cho dự toán
    ws4 = wb.create_sheet("Tong hop block")
    ws4.append(["Tên block", "Tổng số lượng"])
    gop = defaultdict(int)
    for (_, _, b), n in th.block.items():
        gop[b] += n
    for b, n in sorted(gop.items()):
        ws4.append([b, n])

    for s in wb.worksheets:
        for o in s[1]:
            o.font = Font(bold=True)
        for cot in s.columns:
            s.column_dimensions[cot[0].column_letter].width = max(12, max(len(str(c.value or "")) for c in cot) + 2)
    wb.save(duong_dan)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("nguon", nargs="*", help="File .dxf/.dwg hoặc thư mục")
    ap.add_argument("--com-dang-mo", action="store_true", help="Thống kê bản vẽ đang mở trong AutoCAD")
    ap.add_argument("--loc-layer", help="Chỉ lấy layer có tiền tố này, ví dụ PCCC-")
    ap.add_argument("--don-vi", choices=["mm", "m"], default="mm")
    ap.add_argument("--ra", default="khoi_luong_pccc.xlsx")
    args = ap.parse_args()

    files = []
    for n in args.nguon:
        p = Path(n)
        files += sorted([*p.glob("*.dxf"), *p.glob("*.dwg")]) if p.is_dir() else [p]

    th = TongHop()
    dwg = [f for f in files if f.suffix.lower() == ".dwg"]
    for f in files:
        if f.suffix.lower() == ".dxf":
            doc_dxf(str(f), th, args.loc_layer)
            print(f"Đã đọc {f.name}")
    if dwg or args.com_dang_mo:
        import acad_com as ac
        acad = ac.ket_noi()
        if args.com_dang_mo:
            doc_com(ac.goi_lai(lambda: acad.ActiveDocument), th, args.loc_layer)
            print("Đã đọc bản vẽ đang mở")
        for f in dwg:
            doc = ac.goi_lai(lambda: acad.Documents.Open(str(f.resolve()), True))  # chỉ đọc
            try:
                doc_com(doc, th, args.loc_layer)
                print(f"Đã đọc {f.name}")
            finally:
                ac.goi_lai(lambda: doc.Close(False))

    if not (th.block or th.dai or th.dien_tich):
        print("Không có dữ liệu nào để thống kê.")
        return
    xuat_excel(th, args.ra, 1000.0 if args.don_vi == "mm" else 1.0)
    print(f"Đã xuất {args.ra}: {sum(th.block.values())} block, "
          f"{len(th.dai)} nhóm chiều dài, {len(th.dien_tich)} nhóm diện tích.")


if __name__ == "__main__":
    main()
