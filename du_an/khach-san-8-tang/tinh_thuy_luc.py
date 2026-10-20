"""Tính thuỷ lực sơ bộ khách sạn 8 tầng: sprinkler (từng tầng), họng nước, chọn bơm và bể.

  python tinh_thuy_luc.py <scripts> KS_PCCC.dxf.json THUY_LUC_PCCC.xlsx
"""
import json
import math
import sys

sys.path.insert(0, sys.argv[1])
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import ong as ONG
import rai_thiet_bi as rtb
import thuy_luc as TL

tk = json.load(open(sys.argv[2], encoding="utf-8"))
ts = json.loads((rtb.THU_MUC / "thong_so.json").read_text(encoding="utf-8"))
tl = ts["thuy_luc"]
nhom = ts["sprinkler"]["nhom"]["1"]

# Cao độ sàn theo mặt cắt 1-1 (m)
CAO_DO = {"TẦNG 1": 0.0, "TẦNG 2": 3.3, "TẦNG 3-8": 6.4, "TẦNG 4": 9.5, "TẦNG 5": 12.6, "TẦNG 6": 15.7,
          "TẦNG 7": 18.8, "TẦNG 8": 21.9, "TẦNG TUM": 25.0}
TEN = {"TẦNG 3-8": "Tầng 3"}
L_CAP_SPR = tk["ong_spr"]["TẦNG 1"].get("DN100", 0)   # ống cấp ngang từ trạm bơm (m)
L_CAP_HONG = tk["ong_hong"]["TẦNG 1"].get("DN100", 0)

# ---------- Sprinkler từng tầng ----------
ket_qua = []
for tang, mang in tk["mang_spr"].items():
    kq = TL.tinh_tang_spr(mang, nhom, tl, ONG.dn_spr)
    H, tp = TL.ap_luc_bom_spr(kq, CAO_DO[tang], L_CAP_SPR, tl)
    ket_qua.append((tang, kq, H, tp))
    print(f"{TEN.get(tang, tang.title()):9}: {kq['so_dau_hd']:2} đầu hoạt động, {kq['dien_tich_hd_m2']:5.1f} m², "
          f"Q = {kq['Q_lpm']:6.1f} l/ph ({kq['Q_lpm'] / 60:4.2f} l/s), P chân trục = {kq['P_bar']:.2f} bar, "
          f"H bơm = {H:.2f} bar")
tang_bl, kq_bl, H_spr, tp_spr = max(ket_qua, key=lambda r: r[2])
Q_spr = max(r[1]["Q_lpm"] for r in ket_qua)

# ---------- Họng nước ----------
L_truc_hong = CAO_DO["TẦNG 8"] + tl["hong"]["tam_van_tren_san_m"]
Q_hong, H_hong, tp_hong = TL.tinh_hong(CAO_DO["TẦNG 8"], L_truc_hong, L_CAP_HONG, tl)
print(f"Họng nước: Q = {Q_hong:.0f} l/ph, H bơm = {H_hong:.2f} bar")

# ---------- Bơm & bể ----------
du_tru = tl["bom"]["he_so_du_tru_cot_ap"]
Q_bom = Q_spr + Q_hong
H_bom = max(H_spr, H_hong) * du_tru
t_spr = nhom["thoi_gian_phut"]
t_hong = tl["hong"]["thoi_gian_gio"]
V_spr = Q_spr / 1000 * t_spr
V_hong = Q_hong / 1000 * t_hong * 60
V = V_spr + V_hong
V_de_xuat = math.ceil(V * 1.1)
print(f"Bơm chính: Q = {Q_bom / 1000 * 60:.1f} m³/h ({Q_bom / 60:.2f} l/s), H = {H_bom * 10.2:.1f} m")
print(f"Bể: sprinkler {V_spr:.1f} m³ + họng {V_hong:.1f} m³ = {V:.1f} m³ -> đề xuất {V_de_xuat} m³")

# ---------- Excel ----------
wb = Workbook()
dam, nen = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")


def td(ws, hang):
    ws.append(hang)
    for o in ws[ws.max_row]:
        o.font, o.fill = dam, nen
        o.alignment = Alignment(wrap_text=True, vertical="center")


ws = wb.active
ws.title = "Tong hop"
td(ws, ["Hạng mục", "Giá trị", "Đơn vị", "Ghi chú"])
dong = [
    ("SPRINKLER (nhóm 1, TCVN 7336:2021)", None, None, None),
    ("Cường độ phun", nhom["cuong_do_l_s_m2"], "l/s.m²", "Bảng thông số theo nhóm"),
    ("Diện tích tính toán", nhom["dien_tich_tinh_toan_m2"], "m²", ""),
    ("Lưu lượng tối thiểu = cường độ × diện tích", round(kq_bl["Q_toi_thieu_lpm"] / 60, 2), "l/s", ""),
    ("Tầng bất lợi nhất", TEN.get(tang_bl, tang_bl.title()), "", "Áp lực bơm cần lớn nhất"),
    ("Số đầu phun hoạt động trong vùng tính toán", kq_bl["so_dau_hd"], "đầu", f"{kq_bl['dien_tich_hd_m2']:.1f} m²"),
    ("Lưu lượng đầu bất lợi nhất", round(kq_bl["dau_bat_loi"]["q_lpm"], 1), "l/phút",
     f"bảo vệ {kq_bl['dau_bat_loi']['dt_bao_ve_m2']} m², cách trục {kq_bl['dau_bat_loi']['chieu_dai_ong_toi_truc_m']:.1f} m ống"),
    ("Áp lực đầu bất lợi nhất", round(kq_bl["dau_bat_loi"]["P_bar"], 2), "bar", f"K = {tl['sprinkler']['K_l_phut_bar05']}"),
    ("Lưu lượng hệ sprinkler (sau cân bằng)", round(Q_spr / 60, 2), "l/s", f"{Q_spr * 60 / 1000:.1f} m³/h, lấy lớn nhất các tầng ({TEN.get(max(ket_qua, key=lambda r: r[1]['Q_lpm'])[0], max(ket_qua, key=lambda r: r[1]['Q_lpm'])[0].title())})"),
    ("Áp lực cần tại đầu đẩy bơm (sprinkler)", round(H_spr, 2), "bar", f"{H_spr * 10.2:.1f} m"),
    ("HỌNG NƯỚC VÁCH TƯỜNG", None, None, None),
    ("Số cột nước × lưu lượng mỗi cột", f"{tl['hong']['so_cot_nuoc']} × {tl['hong']['q_moi_cot_l_s']}", "l/s",
     "Giả thiết, kiểm tra theo QCVN 06:2022 cho F1.2"),
    ("Áp lực cần tại đầu đẩy bơm (họng)", round(H_hong, 2), "bar", f"{H_hong * 10.2:.1f} m"),
    ("BƠM CHỮA CHÁY (bơm chung, chạy đồng thời 2 hệ)", None, None, None),
    ("Lưu lượng bơm chính", round(Q_bom * 60 / 1000, 1), "m³/h", f"{Q_bom / 60:.2f} l/s"),
    ("Cột áp bơm chính", round(H_bom * 10.2, 1), "m", f"lấy max 2 hệ × {du_tru} dự trữ"),
    ("Bơm dự phòng", "1 bơm cùng thông số", "", "Diesel hoặc điện có nguồn dự phòng"),
    ("Bơm bù áp", f"~1–2 m³/h, {(H_bom + tl['bom']['ap_bu_bom_bu_bar']) * 10.2:.0f} m", "", "Duy trì áp mạng"),
    ("BỂ NƯỚC CHỮA CHÁY", None, None, None),
    ("Dự trữ sprinkler", round(V_spr, 1), "m³", f"{Q_spr / 60:.2f} l/s × {t_spr} phút"),
    ("Dự trữ họng nước", round(V_hong, 1), "m³", f"{Q_hong / 60:.2f} l/s × {t_hong} giờ"),
    ("Tổng dung tích tính toán", round(V, 1), "m³", "Chưa gồm nước chữa cháy ngoài nhà"),
    ("Dung tích đề xuất", V_de_xuat, "m³", "+10%, làm tròn"),
]
for d in dong:
    ws.append(list(d))
    if d[1] is None:
        ws[ws.max_row][0].font = dam
ws.append([])
ws.append(["Kết quả sơ bộ, dùng để chọn thiết bị ban đầu. Phải tính thuỷ lực đầy đủ theo TCVN 7336:2021 và "
           "kiểm tra số liệu QCVN 06:2022 trước khi lập hồ sơ thẩm duyệt."])

ws = wb.create_sheet("Sprinkler tung tang")
td(ws, ["Tầng", "Cao độ sàn (m)", "Số đầu HĐ", "Diện tích HĐ (m²)", "Q (l/s)", "P chân trục tầng (bar)",
        "H đầu đẩy bơm (bar)", "H (m)"])
for tang, kq, H, _ in ket_qua:
    ws.append([TEN.get(tang, tang.title()), CAO_DO[tang], kq["so_dau_hd"], round(kq["dien_tich_hd_m2"], 1),
               round(kq["Q_lpm"] / 60, 2), round(kq["P_bar"], 2), round(H, 2), round(H * 10.2, 1)])

ws = wb.create_sheet("Ap luc bom")
td(ws, ["Hệ", "Thành phần", "bar"])
for he, tp, tong in [("Sprinkler (" + TEN.get(tang_bl, tang_bl.title()) + ")", tp_spr, H_spr),
                     ("Họng nước", tp_hong, H_hong)]:
    for k, v in tp.items():
        ws.append([he, k, round(v, 3)])
    ws.append([he, "TỔNG", round(tong, 2)])
    ws[ws.max_row][1].font = ws[ws.max_row][2].font = dam

ws = wb.create_sheet("Gia thiet")
td(ws, ["Thông số", "Giá trị"])
for nhom_ts, d in tl.items():
    if nhom_ts.startswith("_"):
        continue
    for k, v in d.items():
        if not k.startswith("_"):
            ws.append([f"{nhom_ts}.{k}", json.dumps(v, ensure_ascii=False) if isinstance(v, dict) else v])

for s in wb.worksheets:
    for cot in s.columns:
        s.column_dimensions[cot[0].column_letter].width = min(70, max(12, max(len(str(c.value or "")) for c in cot[1:]) + 2))
wb.save(sys.argv[3])
print("Đã xuất", sys.argv[3])
