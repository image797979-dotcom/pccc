"""Bảng khối lượng PCCC theo tầng (tránh đếm trùng tờ 'tầng 3-8' với tờ từng tầng)."""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

tk = json.load(open(sys.argv[1], encoding="utf-8"))
wb = Workbook()
dam = Font(bold=True)
nen = PatternFill("solid", fgColor="DDEBF7")


def tieu_de(ws, cot):
    ws.append(cot)
    for o in ws[ws.max_row]:
        o.font, o.fill = dam, nen
        o.alignment = Alignment(wrap_text=True, vertical="center")


# Sprinkler & họng/bình: có tờ riêng T1, T2, T3-8 (dùng cho T3), T4..T8, tum
TANG_RIENG = [("Tầng 1", "TẦNG 1"), ("Tầng 2", "TẦNG 2"), ("Tầng 3", "TẦNG 3-8")] + \
             [(f"Tầng {i}", f"TẦNG {i}") for i in range(4, 9)] + [("Tầng tum", "TẦNG TUM")]

ws = wb.active
ws.title = "Sprinkler"
tieu_de(ws, ["Tầng", "Tờ bản vẽ", "Số phòng", "Số đầu phun", "Ghi chú"])
for ten, khoa in TANG_RIENG:
    d = tk["sprinkler"][khoa]
    ws.append([ten, khoa, d["so_phong"], d["so_dau"], "; ".join(c for c in d["canh_bao"] if "dải hẹp" not in c)[:250]])
ws.append(["TỔNG", "", "", f"=SUM(D2:D{ws.max_row})", "Nhóm 1 (TCVN 7336:2021), cách tường ≥ 500 mm"])
ws[ws.max_row][0].font = ws[ws.max_row][3].font = dam

ws = wb.create_sheet("Bao chay")
tieu_de(ws, ["Tầng", "Tờ bản vẽ", "Số tầng áp dụng", "Đầu báo khói", "Đầu báo nhiệt", "Nút ấn", "Chuông + đèn",
             "Tủ trung tâm", "Cáp tín hiệu (m, sơ bộ)"])
for ten, khoa, n in [("Tầng 1", "TẦNG 1", 1), ("Tầng 2", "TẦNG 2", 1), ("Tầng 3–8 (điển hình)", "TẦNG 3-8", 6),
                     ("Tầng tum", "TẦNG TUM", 1)]:
    d = tk["bao_chay"][khoa]
    ws.append([ten, khoa, n, d["bao_khoi"], d["bao_nhiet"], d["nut_an"], d["chuong_den"], d["tu_tt"], d["cap_m"]])
r = ws.max_row
ws.append(["TỔNG", "", ""] + [f"=SUMPRODUCT($C2:$C{r},{c}2:{c}{r})" for c in "DEFGHI"])
for o in ws[ws.max_row]:
    o.font = dam

ws = wb.create_sheet("Hong nuoc - Binh")
tieu_de(ws, ["Tầng", "Tờ bản vẽ", "Hộp họng nước", "Bình chữa cháy", "Xa nhất tới họng (m, đường chim bay)",
             "Xa nhất tới bình (m)"])
for ten, khoa in TANG_RIENG:
    d = tk["hong_binh"][khoa]
    ws.append([ten, khoa, d["hong"], d["binh"], d.get("hong_tam_xa_nhat_m", "-"), d["binh_xa_nhat_m"]])
r = ws.max_row
ws.append(["TỔNG", "", f"=SUM(C2:C{r})", f"=SUM(D2:D{r})", "", ""])
for o in ws[ws.max_row]:
    o.font = dam

ws = wb.create_sheet("Ong")
DN = ["DN25", "DN32", "DN40", "DN50", "DN65", "DN80", "DN100"]
tieu_de(ws, ["Hệ", "Tầng"] + [f"{d} (m)" for d in DN] + ["Ghi chú"])
for he, khoa_tk in [("Sprinkler", "ong_spr"), ("Họng nước", "ong_hong")]:
    for ten, khoa in TANG_RIENG:
        d = tk.get(khoa_tk, {}).get(khoa)
        if d:
            ws.append([he, ten] + [d.get(x, 0) for x in DN] + ["Ống ngang trên mặt bằng"])
# Trục đứng: theo cao độ trên mặt cắt (tầng 1 ±0.000, tầng 8 +21.900, tum +25.000)
ws.append(["Sprinkler", "Trục đứng T1 → tum"] + [0] * 6 + [25.0, "Ước tính theo mặt cắt 1-1"])
ws.append(["Họng nước", "Trục đứng T1 → T8"] + [0] * 6 + [23.2, "Tới cao độ hộp họng tầng 8 (+21.900 + 1.250)"])
r = ws.max_row
ws.append(["TỔNG", ""] + [f"=SUM({c}2:{c}{r})" for c in "CDEFGHI"] + ["Chưa tính hao hụt, phụ kiện"])
for o in ws[ws.max_row]:
    o.font = dam
ws.append([])
ws.append(["Đường kính chọn sơ bộ theo số đầu phun phía sau đoạn ống (bảng ống nguy cơ thấp, nhóm 1). "
           "Phải tính thuỷ lực theo TCVN 7336:2021 trước khi thi công."])

for s in wb.worksheets:
    s.row_dimensions[1].height = 32
    for cot in s.columns:
        s.column_dimensions[cot[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in cot[1:]) + 2))
    s.freeze_panes = "A2"
wb.save(sys.argv[2])
print("Đã xuất", sys.argv[2])
