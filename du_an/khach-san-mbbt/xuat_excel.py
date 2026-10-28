"""Xuất khối lượng thiết bị theo tầng và thuỷ lực sơ bộ ra Excel từ thiet_ke.json.

  python xuat_excel.py   ->  KHOI_LUONG_THUY_LUC_PCCC.xlsx
"""
import json
import sys

sys.path.insert(0, r"C:\Users\Administrator\Documents\pccc\.claude\skills\pccc-autocad\scripts")
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import ong as ONG

TK = json.load(open("thiet_ke.json", encoding="utf-8"))
TEN = {"HAM": "Tầng hầm", "TRET": "Tầng trệt", "LUNG": "Tầng lửng", "T2": "Tầng 2", "T3": "Tầng 3", "T4": "Tầng 4",
       "T5": "Tầng 5", "T6": "Tầng 6", "T7": "Tầng 7", "SAN_THUONG": "Sân thượng"}
dam, nen = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
wb = Workbook()


def tieu_de(ws, cot):
    ws.append(cot)
    for o in ws[ws.max_row]:
        o.font, o.fill = dam, nen
        o.alignment = Alignment(wrap_text=True, vertical="center")


def tong(ws, c0=2):
    n = ws.max_row
    ws.append(["TỔNG"] + [f"=SUM({chr(64 + c)}2:{chr(64 + c)}{n})" for c in range(c0, ws.max_column + 1)])
    for o in ws[ws.max_row]:
        o.font = dam


ws = wb.active
ws.title = "Báo cháy"
tieu_de(ws, ["Tầng", "Đầu báo khói", "Đầu báo nhiệt", "Nút nhấn", "Chuông + đèn", "Tủ trung tâm", "Đèn FIRE",
             "Cáp báo cháy (m)"])
for t, T in TK["tang"].items():
    b = T["bc"]
    L = sum(abs(c[1][0] - c[0][0]) + abs(c[2][1] - c[1][1]) for c in T["cap"]) / 1000
    ws.append([TEN[t]] + [len(b.get(k, [])) for k in ("bao_khoi", "bao_nhiet", "nut_an", "chuong_den", "tu_tt",
                                                       "den_fire")] + [round(L * 1.1, 1)])
tong(ws)

ws = wb.create_sheet("Sprinkler")
dns = sorted({ONG.dn_spr(n) for T in TK["tang"].values() for _, _, n, _ in T["ong_spr"]})
tieu_de(ws, ["Tầng", "Đầu phun"] + [f"Ống DN{d} (m)" for d in dns])
for t, T in TK["tang"].items():
    L = ONG.gom_chieu_dai([(tuple(a), tuple(b), n, l) for a, b, n, l in T["ong_spr"]], ONG.dn_spr)
    ws.append([TEN[t], len(T["spr"])] + [round(L.get(d, 0) / 1000, 1) for d in dns])
tong(ws)

ws = wb.create_sheet("Họng + bình")
tieu_de(ws, ["Tầng", "Hộp họng FHR", "Bình bột ABC 8 kg", "Bình CO2 5 kg"])
for t, T in TK["tang"].items():
    ws.append([TEN[t], len(T["hong"]), len(T["binh"]), len(T["binh"])])
tong(ws)

ws = wb.create_sheet("Chiếu sáng sự cố")
ks = ["exit", "exit_huong", "den_su_co", "sdtn", "sdtnp"]
tieu_de(ws, ["Tầng", "Đèn EXIT", "EXIT chỉ hướng", "Đèn sự cố", "Sơ đồ thoát nạn tầng", "Sơ đồ thoát nạn phòng"])
for t, T in TK["tang"].items():
    ws.append([TEN[t]] + [len(T["cs"].get(k, [])) for k in ks])
tong(ws)

ws = wb.create_sheet("Thuỷ lực sơ bộ")
tl = TK["thuy_luc"]
tieu_de(ws, ["Hạng mục", "Giá trị", "Ghi chú"])
for dong in [
    ("Nhóm nguy hiểm cháy", "1", "Người dùng chốt cho toàn công trình (TCVN 7336:2021)"),
    ("Cường độ phun / diện tích tính toán", "0,08 l/s.m² / 60 m²", "thong_so.json"),
    ("Tầng cần lưu lượng lớn nhất", TEN[tl["tang_Q_max"]], "vùng tính toán 60 m² ở phía xa trục đứng"),
    ("Tầng cần áp lực lớn nhất", TEN[tl["tang_bat_loi"]], "tầng cao nhất có sprinkler đủ diện tích"),
    ("Lưu lượng sprinkler (l/s)", tl["Q_spr_l_s"], "phòng nhỏ nên nhiều đầu hoạt động đồng thời"),
    ("Áp lực bơm cho sprinkler (bar)", tl["P_spr_bar"], ""),
    ("Lưu lượng họng nước (l/s)", tl["Q_hong_l_s"], "1 họng × 2,5 l/s"),
    ("Áp lực bơm cho họng (bar)", tl["P_hong_bar"], "áp tại van họng cao nhất 2,0 bar"),
    ("LƯU LƯỢNG BƠM CHỌN (l/s)", tl["Q_bom_l_s"], "Q sprinkler + Q họng"),
    ("CỘT ÁP BƠM CHỌN (m)", tl["H_bom_m"], "max(sprinkler, họng) × 1,1"),
    ("THỂ TÍCH BỂ (m³)", tl["V_be_chon_m3"], "Q spr × 30 phút + Q họng × 3 giờ; chưa gồm chữa cháy ngoài nhà"),
]:
    ws.append(list(dong))
ws.append([])
tieu_de(ws, ["Tầng", "Q tầng (l/s)", "Áp lực bơm cần (bar)"])
for t, v in tl["tung_tang"].items():
    ws.append([TEN[t], v["Q_l_s"], v["P_bom_bar"]])
ws.append([])
ws.append(["Thành phần áp lực (tầng bất lợi)", "bar"])
for k, v in tl["thanh_phan_spr"].items():
    ws.append([k, v])

ws = wb.create_sheet("Cảnh báo")
tieu_de(ws, ["Nội dung"])
for c in TK["canh_bao"]:
    ws.append([c])
ws.append(["Tầng hầm: khối kín không cửa phía bắc (≈42,6 m², có ký hiệu dốc) không bố trí thiết bị – kiểm tra công năng."])
ws.append(["Tầng 7: sân thượng hở phía đông bắc (rãnh thoát nước) không bố trí đầu báo/sprinkler."])
ws.append(["Không xét dầm, đèn, miệng gió, vật cản trên trần khi rải thiết bị."])

for w in wb.worksheets:
    for col in w.columns:
        w.column_dimensions[col[0].column_letter].width = max(12, min(60, max(len(str(c.value or "")) for c in col) + 2))
wb.save("KHOI_LUONG_THUY_LUC_PCCC.xlsx")
print("đã lưu KHOI_LUONG_THUY_LUC_PCCC.xlsx")
